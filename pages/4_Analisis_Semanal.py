import streamlit as st
import requests
from datetime import datetime, timedelta
import pandas as pd
import numpy as np
from statistics import mean
import base64
import io
from docx import Document
from docx.shared import Pt

# --- CONFIGURACIÓN DE LA PÁGINA ---
st.set_page_config(layout="wide", page_title="Análisis Semanal")

# --- FUNCIONES AUXILIARES ---
def format_duration_hm(seconds):
    """Formatea segundos a un formato Xh Ym o Ym para la tabla final."""
    if pd.isna(seconds) or not isinstance(seconds, (int, float)) or seconds < 0:
        return "0m"
    seconds = int(seconds)
    hours, remainder = divmod(seconds, 3600)
    minutes, _ = divmod(remainder, 60)
    if hours > 0:
        return f"{hours}h {minutes}m"
    return f"{minutes}m"

def dataframe_to_word(df):
    """Convierte un DataFrame de Pandas a un documento de Word en memoria."""
    doc = Document()
    doc.add_heading('Resumen Semanal de Entrenamiento', level=1)
    
    table = doc.add_table(rows=1, cols=df.shape[1])
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    for j, col_name in enumerate(df.columns):
        hdr_cells[j].text = str(col_name)
        hdr_cells[j].paragraphs[0].runs[0].font.bold = True

    for index, row in df.iterrows():
        row_cells = table.add_row().cells
        for j, col_name in enumerate(df.columns):
            row_cells[j].text = str(row[col_name])

    mem_file = io.BytesIO()
    doc.save(mem_file)
    mem_file.seek(0)
    return mem_file

# --- FUNCIONES DE OBTENCIÓN Y PROCESAMIENTO DE DATOS ---
@st.cache_data(ttl=3600)
def fetch_data(url, headers):
    try:
        response = requests.get(url, headers=headers)
        response.raise_for_status()
        return response.json()
    except requests.exceptions.RequestException as e:
        st.error(f"Error de conexión con la API: {e}")
        return []

def process_wellness_data(data, date_obj_for_sunday_check):
    if not data: return {}
    rhr_list = [d.get('restingHR') for d in data if d.get('restingHR') is not None]
    hrv_list = [d.get('hrv') for d in data if d.get('hrv') is not None]
    bb_max_list = [d.get('BodyBatteryMax') for d in data if d.get('BodyBatteryMax') is not None]
    bb_min_list = [d.get('BodyBatteryMin') for d in data if d.get('BodyBatteryMin', 0) > 0]
    sleep_list = [d.get('sleepScore') for d in data if d.get('sleepScore') is not None]
    sunday_data = next((d for d in reversed(data) if d.get('id') and datetime.strptime(d['id'], '%Y-%m-%d').date() == date_obj_for_sunday_check), None)
    sunday_ctl = sunday_data.get('ctl') if sunday_data else None
    sunday_atl = sunday_data.get('atl') if sunday_data else None
    sunday_tsb = (sunday_ctl - sunday_atl) if sunday_ctl is not None and sunday_atl is not None else None
    return {
        'RHR_Avg': mean(rhr_list) if rhr_list else None, 'HRV_Avg': mean(hrv_list) if hrv_list else None,
        'BodyBatteryMax_Avg': mean(bb_max_list) if bb_max_list else None,
        'BodyBatteryMin_Avg': mean(bb_min_list) if bb_min_list else None,
        'SleepScore_Avg': mean(sleep_list) if sleep_list else None, 'CTL_Sunday': sunday_ctl,
        'ATL_Sunday': sunday_atl, 'TSB_Sunday': sunday_tsb,
    }

def calculate_training_metrics(activities_data):
    if not activities_data: return {}
    total_tss = 0
    total_hr_zone_times, total_power_zone_times = [0]*7, [0]*7
    efficiency_list, power_hr_list, power_hr_z2_list = [], [], []
    for activity in activities_data:
        tss = activity.get('icu_training_load', 0) or 0
        if activity.get('type') == 'WeightTraining': tss = 10
        total_tss += tss
        power_norm = activity.get("icu_weighted_avg_watts")
        hr_avg = activity.get("average_heartrate")
        power_avg = activity.get("icu_average_watts")
        power_hr_z2 = activity.get("icu_power_hr_z2")
        if hr_avg and hr_avg > 0:
            if power_norm and power_norm > 0: efficiency_list.append(power_norm / hr_avg)
            if power_avg and power_avg > 0: power_hr_list.append(power_avg / hr_avg)
        if power_hr_z2 and power_hr_z2 > 0: power_hr_z2_list.append(power_hr_z2)
        hr_zones = activity.get('icu_hr_zone_times')
        if hr_zones:
            for i, secs in enumerate(hr_zones):
                if i < 7: total_hr_zone_times[i] += (secs or 0)
        power_zones = activity.get('icu_zone_times')
        if power_zones:
            for zone_data in power_zones:
                if 'id' in zone_data and isinstance(zone_data['id'], str) and zone_data['id'].startswith('Z'):
                    try:
                        zone_index = int(zone_data['id'][1:]) - 1
                        if 0 <= zone_index < 7: total_power_zone_times[zone_index] += (zone_data.get('secs', 0) or 0)
                    except (ValueError, IndexError): continue
    return {
        'TSS_Realizado': total_tss, 'HR_Zone_Times': total_hr_zone_times,
        'Power_Zone_Times': total_power_zone_times, 'Eficiencia_Avg': mean(efficiency_list) if efficiency_list else None,
        'Potencia_FC_Avg': mean(power_hr_list) if power_hr_list else None,
        'Potencia_FC_Z2_Avg': mean(power_hr_z2_list) if power_hr_z2_list else None
    }

def get_weekly_analysis(end_date, planned_tss_list):
    api_key = "27i9azt55smmhvg1ogc5gmn7x"
    athlete_id = "i10474"
    base_url = "https://intervals.icu/api/v1"
    headers = {"Authorization": f"Basic {base64.b64encode(f'API_KEY:{api_key}'.encode()).decode()}"}
    weekly_data = []
    for i in range(4):
        week_end_obj = end_date - timedelta(days=7 * i)
        week_start_obj = week_end_obj - timedelta(days=6)
        week_start_str, week_end_str = week_start_obj.strftime("%Y-%m-%d"), week_end_obj.strftime("%Y-%m-%d")
        wellness_data = fetch_data(f"{base_url}/athlete/{athlete_id}/wellness?oldest={week_start_str}&newest={week_end_str}", headers)
        activities_data = fetch_data(f"{base_url}/athlete/{athlete_id}/activities?oldest={week_start_str}&newest={week_end_str}", headers)
        wellness_metrics = process_wellness_data(wellness_data, week_end_obj)
        training_metrics = calculate_training_metrics(activities_data)
        week_summary = {
            "Semana": f"{week_start_obj.strftime('%d/%m')} - {week_end_obj.strftime('%d/%m')}",
            "TSS_Programado": planned_tss_list[i], **training_metrics, **wellness_metrics
        }
        weekly_data.append(week_summary)
    df = pd.DataFrame(weekly_data).set_index("Semana")
    return df.iloc[::-1]

# --- INTERFAZ DE USUARIO CON STREAMLIT ---
st.title("🔬 Análisis Semanal de Entrenamiento")
st.caption("Compara la carga, el rendimiento y la recuperación de las últimas 4 semanas.")
with st.form(key='analysis_form'):
    st.write("**Paso 1: Configura tu análisis**")
    col1, col2 = st.columns([1, 3])
    with col1:
        end_date = st.date_input("Fecha final del análisis (Domingo)", datetime.now().date())
    with col2:
        st.write("**TSS Programado para cada semana**")
        cols_tss = st.columns(4)
        planned_tss = [
            cols_tss[0].number_input("Semana -3", value=330, step=5),
            cols_tss[1].number_input("Semana -2", value=350, step=5),
            cols_tss[2].number_input("Semana -1", value=385, step=5),
            cols_tss[3].number_input("Semana Actual", value=195, step=5)
        ]
    submit_button = st.form_submit_button(label='🚀 Generar Análisis')

if submit_button:
    df_analysis = get_weekly_analysis(end_date, planned_tss[::-1])
    if not df_analysis.empty:
        st.markdown("---")
        st.header("🗓️ Análisis Detallado por Semana")
        tab_names = [f"Semana Actual ({df_analysis.index[0]})"] + [f"Semana {-i} ({df_analysis.index[i]})" for i in range(1, len(df_analysis))]
        tabs = st.tabs(tab_names)
        for i, tab in enumerate(tabs):
            with tab:
                week_data = df_analysis.iloc[i]
                st.subheader("🎯 Cumplimiento de Carga (TSS)")
                tss_real = week_data.get('TSS_Realizado', 0)
                tss_prog = week_data.get('TSS_Programado', 1)
                percentage_completed = (tss_real / tss_prog * 100) if tss_prog > 0 else 0
                progress_bar_value = min(1.0, tss_real / tss_prog) if tss_prog > 0 else 0
                st.progress(progress_bar_value)
                st.caption(f"**{tss_real:.0f}** TSS realizados de **{tss_prog:.0f}** programados (**{percentage_completed:.0f}%**)")
                st.subheader("📊 Métricas Clave de la Semana")
                kpi_cols = st.columns(3)
                with kpi_cols[0]: st.metric(label="TSB (Frescura) del Domingo", value=f"{week_data.get('TSB_Sunday', 0):.1f}")
                with kpi_cols[1]: st.metric(label="HRV Promedio Semanal", value=f"{week_data.get('HRV_Avg', 0):.1f} ms")
                with kpi_cols[2]: st.metric(label="Sueño Promedio Semanal", value=f"{week_data.get('SleepScore_Avg', 0):.1f}")
                st.subheader("📈 Distribución de Zonas e Insights")
                expander_cols = st.columns(3)
                with expander_cols[0]:
                    with st.expander("Zonas de Potencia"):
                        power_zones = week_data.get('Power_Zone_Times', [0]*7)
                        if sum(power_zones) > 0:
                            df_power = pd.DataFrame({'Minutos': [s / 60 for s in power_zones], 'Zona': [f"Z{i+1}" for i in range(7)]}).set_index('Zona')
                            st.bar_chart(df_power['Minutos'])
                        else: st.caption("Sin datos.")
                with expander_cols[1]:
                    with st.expander("Zonas de Frecuencia Cardíaca"):
                        hr_zones = week_data.get('HR_Zone_Times', [0]*7)
                        if sum(hr_zones) > 0:
                            df_hr = pd.DataFrame({'Minutos': [s / 60 for s in hr_zones], 'Zona': [f"Z{i+1}" for i in range(7)]}).set_index('Zona')
                            st.bar_chart(df_hr['Minutos'])
                        else: st.caption("Sin datos.")
                with expander_cols[2]:
                    with st.expander("💡 Insights del Coach"):
                        ratio = (week_data.get('ATL_Sunday', 0) / week_data.get('CTL_Sunday', 1)) if week_data.get('CTL_Sunday') else 0
                        if ratio > 1.3: st.warning(f"**Ratio ATL/CTL: {ratio:.2f}**. Carga muy alta.")
                        elif ratio > 1.1: st.info(f"**Ratio ATL/CTL: {ratio:.2f}**. Carga productiva.")
                        else: st.success(f"**Ratio ATL/CTL: {ratio:.2f}**. Carga de recuperación.")
                        ef = week_data.get('Eficiencia_Avg')
                        if ef: st.info(f"**Eficiencia (NP/FC): {ef:.2f}**.")
                        sleep = week_data.get('SleepScore_Avg')
                        if sleep and sleep < 75: st.warning(f"**Sueño: {sleep:.1f}**. Descanso insuficiente.")
        
        st.markdown("---")
        with st.expander("📋 Ver Tabla Resumen de Datos Completa"):
            df_display = df_analysis.copy()
            if 'Power_Zone_Times' in df_display.columns:
                df_display['Power_Zone_Times'] = df_display['Power_Zone_Times'].apply(
                    lambda zones: [format_duration_hm(s) for s in zones] if isinstance(zones, list) else zones
                )
            if 'HR_Zone_Times' in df_display.columns:
                df_display['HR_Zone_Times'] = df_display['HR_Zone_Times'].apply(
                    lambda zones: [format_duration_hm(s) for s in zones] if isinstance(zones, list) else zones
                )
            
            numeric_cols = df_display.select_dtypes(include=np.number).columns
            st.dataframe(df_display.style.format("{:.1f}", subset=numeric_cols, na_rep="N/A"))
            
            # --- INICIO: LÓGICA DE TRANSPOSICIÓN Y DESCARGA ---
            # 1. Preparar el DataFrame para la transposición, formateando primero los números
            df_for_word = df_display.copy()
            for col in numeric_cols:
                df_for_word[col] = df_for_word[col].apply(lambda x: f"{x:.1f}" if pd.notna(x) else "N/A")

            # 2. Transponer el DataFrame
            df_transposed = df_for_word.reset_index().set_index('Semana').T.reset_index()
            df_transposed = df_transposed.rename(columns={'index': 'Métrica'})

            # 3. Generar el fichero Word
            word_file = dataframe_to_word(df_transposed)
            
            # 4. Crear el botón de descarga
            st.download_button(
                label="📥 Descargar como Word (.docx)",
                data=word_file,
                file_name=f"resumen_semanal_{end_date.strftime('%Y-%m-%d')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            # --- FIN: LÓGICA DE TRANSPOSICIÓN Y DESCARGA ---

    else:
        st.warning("No se pudieron obtener datos para el periodo seleccionado. Revisa la configuración o el rango de fechas.")